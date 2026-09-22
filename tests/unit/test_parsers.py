from pathlib import Path

import pymupdf
from docx import Document as DocxDocument

from app.domain.models import Confidentiality, DocumentRecord, DocumentType, IngestionErrorCode
from app.rag.ingestion.discovery import FileDiscovery
from app.rag.ingestion.manifest import ManifestStatus, ManifestStore
from app.rag.ingestion.service import ParserService
from app.rag.parsing.docx import DocxParser
from app.rag.parsing.registry import default_parser_registry
from app.rag.parsing.text import MarkdownParser, TextParser


def make_document(discovered) -> DocumentRecord:
    return DocumentRecord(
        document_id=discovered.document_id,
        document_version=discovered.document_version,
        title=Path(discovered.relative_path).stem,
        source_uri=discovered.relative_path,
        document_type=DocumentType.OTHER,
        confidentiality=Confidentiality.SYNTHETIC,
        content_checksum=discovered.content_checksum,
    )


def test_markdown_parser_preserves_order_sections_and_source_lines(tmp_path: Path) -> None:
    path = tmp_path / "guide.md"
    path.write_text("# Overview\n\nIntro text.\n\n- first\n- second\n", encoding="utf-8")
    discovered = FileDiscovery(tmp_path, max_file_size_bytes=1024).discover().files[0]

    blocks = MarkdownParser().parse(path, document=make_document(discovered))

    assert [block.block_type.value for block in blocks] == ["heading", "paragraph", "list"]
    assert blocks[0].section_path == ("Overview",)
    assert blocks[1].source_location.line_start == 3
    assert blocks[2].source_location.line_end == 6
    assert [block.order for block in blocks] == [0, 1, 2]


def test_text_parser_identifies_paragraph_and_list_blocks(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("First line.\nSecond line.\n\n1. one\n2. two\n", encoding="utf-8")
    discovered = FileDiscovery(tmp_path, max_file_size_bytes=1024).discover().files[0]

    blocks = TextParser().parse(path, document=make_document(discovered))

    assert [block.block_type.value for block in blocks] == ["paragraph", "list"]
    assert blocks[0].text == "First line.\nSecond line."
    assert blocks[1].source_location.line_start == 4


def test_docx_parser_preserves_heading_paragraph_table_order(tmp_path: Path) -> None:
    path = tmp_path / "guide.docx"
    source = DocxDocument()
    source.add_heading("Overview", level=1)
    source.add_paragraph("A short explanation.")
    table = source.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Key"
    table.rows[0].cells[1].text = "Value"
    source.save(path)
    discovered = FileDiscovery(tmp_path, max_file_size_bytes=100000).discover().files[0]

    blocks = DocxParser().parse(path, document=make_document(discovered))

    assert [block.block_type.value for block in blocks] == ["heading", "paragraph", "table"]
    assert blocks[0].section_path == ("Overview",)
    assert blocks[2].text == "Key | Value"


def test_pdf_parser_preserves_page_provenance(tmp_path: Path) -> None:
    path = tmp_path / "guide.pdf"
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "page one")
    pdf.new_page().insert_text((72, 72), "page two")
    pdf.save(path)
    pdf.close()
    discovered = FileDiscovery(tmp_path, max_file_size_bytes=100000).discover().files[0]
    registry = default_parser_registry()

    blocks = registry.for_path(path).parse(path, document=make_document(discovered))

    assert [block.text for block in blocks] == ["page one", "page two"]
    assert blocks[1].source_location.page_start == 2


def test_parser_service_isolates_corrupt_pdf_and_updates_manifest(tmp_path: Path) -> None:
    (tmp_path / "valid.txt").write_text("valid content", encoding="utf-8")
    (tmp_path / "broken.pdf").write_bytes(b"not a PDF")
    discovery = FileDiscovery(tmp_path, max_file_size_bytes=1024).discover()
    manifest_store = ManifestStore(tmp_path / "manifest.jsonl")
    service = ParserService(default_parser_registry(), manifest_store=manifest_store)

    outcomes = service.parse_batch(tuple((item, make_document(item)) for item in discovery.files))

    successful = [outcome for outcome in outcomes if outcome.error is None]
    failed = [outcome for outcome in outcomes if outcome.error is not None]
    assert len(successful) == 1
    assert len(failed) == 1
    assert failed[0].error.code is IngestionErrorCode.PARSE_FAILED
    assert manifest_store.get(successful[0].document.document_id).status is ManifestStatus.PARSED
    assert manifest_store.get(failed[0].document.document_id).status is ManifestStatus.FAILED
